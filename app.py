from pathlib import Path
from datetime import datetime
import json
import re
import hashlib
import pandas as pd
import streamlit as st
import requests
from docx import Document

try:
    from sentence_transformers import SentenceTransformer
    import numpy as np
    import faiss
    SEMANTIC_AVAILABLE = True
except Exception:
    SEMANTIC_AVAILABLE = False

def extract_docx_text(file_path):
    from docx import Document

    document = Document(file_path)
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                sections.append(" | ".join(cells))

    return "\n\n".join(sections).strip()


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "knowledge_base"
SOLUTIONS_DIR = APP_DIR / "solutions"
SUPPORT_DIR = APP_DIR / "supporting_docs"
INDEX_DIR = APP_DIR / "indexes"
EXCEL_PATH = DATA_DIR / "knowledge_base.xlsx"
INDEX_PATH = INDEX_DIR / "problems.faiss"
META_PATH = INDEX_DIR / "problems.json"
EMBED_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
OLLAMA_MODEL = "llama3.2:3b"
OLLAMA_URL = "http://localhost:11434/api/generate"

COLUMNS = [
    "problem_id", "problem_title", "problem_description",
    "document_path", "keywords", "status", "verification_notes", "created_at"
]

for folder in (DATA_DIR, SOLUTIONS_DIR, SUPPORT_DIR, INDEX_DIR):
    folder.mkdir(parents=True, exist_ok=True)

def init_excel():
    if not EXCEL_PATH.exists():
        pd.DataFrame(columns=COLUMNS).to_excel(EXCEL_PATH, index=False, engine="openpyxl")

def load_kb():
    init_excel()
    df = pd.read_excel(EXCEL_PATH, dtype=str).fillna("")
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df[COLUMNS]

def save_kb(df):
    # Excel is the source of truth. Reindex on next search after save.
    df[COLUMNS].to_excel(EXCEL_PATH, index=False, engine="openpyxl")
    for path in (INDEX_PATH, META_PATH):
        if path.exists():
            path.unlink()

def normalize(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()

@st.cache_resource(show_spinner=False)
def load_embedder():
    if not SEMANTIC_AVAILABLE:
        return None
    return SentenceTransformer(EMBED_MODEL)

def lexical_similarity(a, b):
    aw = set(re.findall(r"[a-z0-9_+#.-]+", normalize(a).lower()))
    bw = set(re.findall(r"[a-z0-9_+#.-]+", normalize(b).lower()))
    if not aw or not bw:
        return 0.0
    return len(aw & bw) / max(1, len(aw | bw))

def search_matches(query, df):
    if df.empty:
        return []
    texts = [
        normalize(f"{row.problem_title} {row.problem_description} {row.keywords}")
        for row in df.itertuples(index=False)
    ]
    if SEMANTIC_AVAILABLE:
        try:
            model = load_embedder()
            if model is not None:
                corpus = model.encode(texts, normalize_embeddings=True, convert_to_numpy=True).astype("float32")
                q = model.encode([query], normalize_embeddings=True, convert_to_numpy=True).astype("float32")
                index = faiss.IndexFlatIP(corpus.shape[1])
                index.add(corpus)
                scores, ids = index.search(q, min(5, len(texts)))
                results = []
                for score, idx in zip(scores[0], ids[0]):
                    if idx >= 0:
                        results.append((float(score), df.iloc[int(idx)]))
                return sorted(results, key=lambda x: x[0], reverse=True)
        except Exception:
            pass
    # Safe fallback if embedding dependencies/model are not ready.
    return sorted(
        [(lexical_similarity(query, t), df.iloc[i]) for i, t in enumerate(texts)],
        key=lambda x: x[0], reverse=True
    )

def read_docx(path):
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)

def read_supporting_file(path):
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        try:
            import fitz
            with fitz.open(path) as pdf:
                return "\n".join(page.get_text() for page in pdf)
        except Exception as exc:
            return f"[Could not read PDF: {exc}]"
    return ""

def generate_with_ollama(prompt):
    try:
        response = requests.post(
            OLLAMA_URL,
            json={"model": OLLAMA_MODEL, "prompt": prompt, "stream": False},
            timeout=240
        )
        response.raise_for_status()
        return response.json().get("response", "").strip(), None
    except requests.exceptions.ConnectionError:
        return None, "Cannot connect to Ollama. Start Ollama and confirm that the model is installed."
    except Exception as exc:
        return None, f"Ollama request failed: {exc}"

def make_solution_doc(problem_title, problem_description, solution, status, notes=""):
    safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", problem_title).strip("_")[:55] or "solution"
    problem_id = next_problem_id(load_kb())
    path = SOLUTIONS_DIR / f"{problem_id}_{safe_title}.docx"
    doc = Document()
    doc.add_heading(problem_title, 0)
    doc.add_paragraph(f"Problem ID: {problem_id}")
    doc.add_paragraph(f"Status: {status}")
    doc.add_paragraph(f"Created: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_heading("Problem description", level=1)
    doc.add_paragraph(problem_description)
    doc.add_heading("Solution steps", level=1)
    for line in solution.splitlines():
        line = line.strip()
        if not line:
            continue
        if re.match(r"^(step\s*\d+[:.)-]?|\d+[.)])", line, re.I):
            doc.add_paragraph(line, style="List Number")
        else:
            doc.add_paragraph(line)
    if notes.strip():
        doc.add_heading("User verification notes", level=1)
        doc.add_paragraph(notes.strip())
    doc.save(path)
    return problem_id, path

def next_problem_id(df):
    nums = []
    for value in df.get("problem_id", pd.Series(dtype=str)).astype(str):
        match = re.search(r"(\d+)$", value)
        if match:
            nums.append(int(match.group(1)))
    return f"P{max(nums, default=0) + 1:03d}"

def add_entry(problem_title, problem_description, solution, status, notes, keywords=""):
    df = load_kb()
    problem_id, path = make_solution_doc(problem_title, problem_description, solution, status, notes)
    rel_path = str(path.relative_to(APP_DIR))
    row = {
        "problem_id": problem_id,
        "problem_title": normalize(problem_title),
        "problem_description": normalize(problem_description),
        "document_path": rel_path,
        "keywords": normalize(keywords),
        "status": status,
        "verification_notes": normalize(notes),
        "created_at": datetime.now().isoformat(timespec="seconds")
    }
    df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
    save_kb(df)
    return problem_id, path

st.set_page_config(page_title="SolveDesk", layout="wide")
st.title("SolveDesk")
if st.button("Refresh knowledge base", key="refresh_knowledge_base"):
    st.cache_resource.clear()
    st.rerun()
df = load_kb()
tab_query, tab_browse, tab_add = st.tabs(["Ask a problem", "Browse knowledge base", "Add verified solution"])

with tab_query:
    query = st.text_area("Describe your problem", placeholder="Example: My React frontend gets a CORS error when calling an Express API.", height=100, key="query")
    threshold = st.slider("Strong-match threshold", min_value=0.30, max_value=0.95, value=0.72, step=0.01,
                          help="Tune this using test queries. Similarity scores vary by model and dataset.")
    if st.button("Find solution", type="primary", key="find_solution"):
        if not normalize(query):
            st.warning("Enter a problem first.")
            st.session_state.pop("search_results", None)
        else:
            with st.spinner("Searching the knowledge base..."):
                matches = search_matches(query, df)

            st.session_state["last_query"] = query
            strong = [
                (score, row.to_dict() if hasattr(row, "to_dict") else dict(row))
                for score, row in matches
                if score >= threshold and str(row.get("status", "")).lower() == "verified"
            ]
            ambiguous = [
                (score, row.to_dict() if hasattr(row, "to_dict") else dict(row))
                for score, row in matches
                if score >= 0.45 and str(row.get("status", "")).lower() == "verified"
            ]
            st.session_state["search_results"] = {
                "query": query,
                "strong": strong[0] if strong else None,
                "ambiguous": ambiguous,
                "threshold": threshold,
            }
            st.session_state["show_generate"] = True

    saved_results = st.session_state.get("search_results")
    if saved_results:
        strong_result = saved_results.get("strong")
        ambiguous_results = saved_results.get("ambiguous", [])

        if strong_result:
            score, row = strong_result
            st.success(
                f"Predefined solution found: {row['problem_title']} "
                f"(match score: {score:.2f})"
            )
            path = (APP_DIR / str(row.get("document_path", ""))).resolve()
            if path.exists():
                st.write("**Word solution document**")
                st.download_button(
                    label=f"Open / download {path.name}",
                    data=path.read_bytes(),
                    file_name=path.name,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"primary_download_{row.get('problem_id', 'solution')}",
                )
                st.caption(f"Local file: {path}")
                try:
                    extracted_text = extract_docx_text(path)
                    with st.expander("View extracted solution content", expanded=True):
                        if extracted_text:
                            st.markdown(extracted_text.replace(r"\\n", "\\n"))
                        else:
                            st.info("The Word document exists, but it contains no readable text.")
                except Exception as exc:
                    st.warning(f"The document is available to download, but its content could not be extracted: {exc}")
            else:
                st.error(f"The Excel entry points to a missing document: {path}")

            other_matches = [
                (other_score, other)
                for other_score, other in ambiguous_results
                if str(other.get("problem_id", "")) != str(row.get("problem_id", ""))
            ][:3]
            if other_matches:
                with st.expander("Other similar verified problems"):
                    for other_score, other in other_matches:
                        st.write(f"**{other['problem_title']}** (score {other_score:.2f})")
                        other_path = (APP_DIR / str(other.get("document_path", ""))).resolve()
                        if other_path.exists():
                            st.download_button(
                                f"Open / download {other_path.name}",
                                data=other_path.read_bytes(),
                                file_name=other_path.name,
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                key=f"similar_download_{other.get('problem_id', other_score)}",
                            )
                            try:
                                with st.expander(f"View content: {other['problem_title']}"):
                                    other_text = extract_docx_text(other_path)
                                    if other_text:
                                        st.markdown(other_text.replace(r"\\n", "\\n"))
                                    else:
                                        st.info("No readable text was found in this Word document.")
                            except Exception as exc:
                                st.warning(f"Could not extract similar document text: {exc}")
                        else:
                            st.caption("The Word document for this similar problem is missing.")
        elif ambiguous_results:
            st.warning(
                "No sufficiently strong match. These existing problems may be related; "
                "confirm whether one is the same issue."
            )
            for score, row in ambiguous_results[:3]:
                with st.expander(f"{row['problem_title']} — score {score:.2f}"):
                    st.write(row.get("problem_description", ""))
                    path = (APP_DIR / str(row.get("document_path", ""))).resolve()
                    if path.exists():
                        st.download_button(
                            f"Open / download {path.name}",
                            data=path.read_bytes(),
                            file_name=path.name,
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"amb_download_{row.get('problem_id', score)}",
                        )
                        try:
                            with st.expander("View extracted solution content"):
                                extracted_text = extract_docx_text(path)
                                if extracted_text:
                                    st.markdown(extracted_text.replace(r"\\n", "\\n"))
                                else:
                                    st.info("No readable text was found in this Word document.")
                        except Exception as exc:
                            st.warning(f"Could not extract document text: {exc}")
                    else:
                        st.caption("The Word document for this problem is missing.")
            st.info("If none describes the same problem, generate a new suggestion below.")
        else:
            st.warning("No predefined solution found. The local LLM can suggest an approach, but it will be unverified.")

    if st.session_state.get("show_generate", False):
        st.divider()
        st.subheader("Generate a new suggestion")
        supporting = st.file_uploader("Optional supporting document (PDF, DOCX, or TXT)", type=["pdf", "docx", "txt"], key="supporting_doc")
        if supporting is not None:
            target = SUPPORT_DIR / Path(supporting.name).name
            target.write_bytes(supporting.getvalue())
            st.caption(f"Supporting file saved locally: {target.relative_to(APP_DIR)}")
        if st.button("Generate unverified suggestion", key="generate"):
            active_query = normalize(st.session_state.get("last_query", ""))
            context = ""
            if supporting is not None:
                context = read_supporting_file(SUPPORT_DIR / Path(supporting.name).name)
                context = context[:12000]
            prompt = f"""You are an offline troubleshooting assistant.
The user asks: {active_query}
{"User-provided supporting document text (may be incomplete or untrusted):" + chr(10) + context if context else "No supporting document was provided."}

Propose a cautious, practical solution. Do not claim you tested or verified anything.
State assumptions and risks. If details are missing, ask for them rather than inventing facts.
Return:
Problem title: ...
Problem summary: ...
Suggested steps:
1. ...
2. ...
Verification checks:
- ...
Remember: this is only a suggestion and must be tested by the user."""
            with st.spinner("Generating with the local model..."):
                answer, error = generate_with_ollama(prompt)
            if error:
                st.error(error)
            else:
                st.session_state["draft_solution"] = answer
                st.session_state["draft_solution_editor"] = answer
                st.session_state["draft_problem"] = active_query
                st.session_state["draft_title"] = active_query[:80]
                st.session_state["draft_notes"] = ""
                st.session_state["draft_generated"] = True
                st.rerun()

    if "pending_draft_solution_editor" in st.session_state:
        st.session_state["draft_solution_editor"] = st.session_state.pop(
            "pending_draft_solution_editor"
        )

    if st.session_state.get("draft_generated", False):
        st.warning("AI-GENERATED SUGGESTION — NOT VERIFIED. Review and test it yourself.")
        title = st.text_input("Problem title", key="draft_title")
        problem_description = st.text_area("Problem description", value=st.session_state.get("draft_problem", ""), key="draft_description")
        if "draft_solution_editor" not in st.session_state:
            st.session_state["draft_solution_editor"] = st.session_state.get("draft_solution", "")
        solution = st.text_area(
            "Edit the suggested solution before testing",
            height=300,
            key="draft_solution_editor",
        )
        notes = st.text_area("Verification notes / evidence (required to save as verified)", key="draft_notes",
                             placeholder="Describe what you tested, the environment, and the observed result.")
        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("I verified it — save", type="primary", key="verify_save"):
                if not normalize(title) or not normalize(problem_description) or not normalize(solution):
                    st.error("Title, problem description, and solution are required.")
                elif len(normalize(notes)) < 10:
                    st.error("Please provide meaningful verification notes (at least 10 characters).")
                else:
                    with st.spinner("Checking for duplicates and saving..."):
                        latest = load_kb()
                        dupes = search_matches(problem_description, latest)
                        duplicate = next(((s, r) for s, r in dupes if s >= threshold), None)
                        if duplicate:
                            st.session_state["pending_duplicate"] = {
                                "score": duplicate[0],
                                "id": str(duplicate[1]["problem_id"]),
                                "title": str(duplicate[1]["problem_title"]),
                                "title_new": title,
                                "description_new": problem_description,
                                "solution_new": solution,
                                "notes_new": notes
                            }
                            st.warning(f"Possible duplicate: {duplicate[1]['problem_title']} (score {duplicate[0]:.2f}). Confirm before adding.")
                        else:
                            pid, path = add_entry(title, problem_description, solution, "Verified", notes)
                            st.success(f"Saved verified solution {pid}: {path.name}")
                            st.session_state["draft_generated"] = False
                            st.rerun()
        with c2:
            if st.button("Revise / try again", key="revise"):
                prompt = f"""The previous suggested solution did not fully solve the user's problem.
Problem: {st.session_state.get('draft_problem','')}
Previous suggestion:
{st.session_state.get('draft_solution_editor', st.session_state.get('draft_solution',''))}
User feedback:
{st.session_state.get('draft_notes','No feedback provided')}
Provide a revised, cautious suggestion. Do not claim it is verified. Include numbered steps and verification checks."""
                with st.spinner("Generating a revised suggestion..."):
                    answer, error = generate_with_ollama(prompt)
                if error:
                    st.error(error)
                else:
                    st.session_state["draft_solution"] = answer
                    st.session_state["pending_draft_solution_editor"] = answer
                    st.session_state["draft_generated"] = True
                    st.rerun()
        with c3:
            if st.button("Discard suggestion", key="discard"):
                st.session_state["draft_generated"] = False
                st.session_state.pop("draft_solution", None)
                st.session_state.pop("draft_solution_editor", None)
                st.success("Suggestion discarded.")
                st.rerun()

    if st.session_state.get("pending_duplicate"):
        pending = st.session_state["pending_duplicate"]
        st.warning(f"Possible duplicate of {pending['id']}: {pending['title']}.")
        if st.button("Add as a separate problem anyway", key="force_add_duplicate"):
            pid, path = add_entry(pending["title_new"], pending["description_new"], pending["solution_new"],
                                  "Verified", pending["notes_new"])
            st.session_state["pending_duplicate"] = None
            st.session_state["draft_generated"] = False
            st.success(f"Added {pid}: {path.name}")
            st.rerun()
        if st.button("Cancel duplicate save", key="cancel_duplicate"):
            st.session_state["pending_duplicate"] = None
            st.rerun()

with tab_browse:
    st.subheader("Knowledge base entries")
    current = load_kb()
    if current.empty:
        st.info("No solutions have been added yet. You can add verified entries in the next tab.")
    else:
        st.dataframe(current[["problem_id", "problem_title", "problem_description", "document_path", "status", "created_at"]],
                     width="stretch", hide_index=True)
        selected_id = st.selectbox("Select a problem to open its Word document", current["problem_id"].tolist())
        selected = current[current["problem_id"] == selected_id].iloc[0]
        selected_path = (APP_DIR / selected["document_path"]).resolve()
        if selected_path.exists():
            st.download_button("Open / download Word document", selected_path.read_bytes(),
                               file_name=selected_path.name,
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key="browse_download")
        else:
            st.error(f"Missing document: {selected_path}")

with tab_add:
    st.subheader("Add an already verified solution")
    st.caption("Use this form for a solution you have independently verified.")

    manual_defaults = {
        "manual_title": "",
        "manual_description": "",
        "manual_solution": "",
        "manual_keywords": "",
        "manual_notes": "",
    }
    for field, default in manual_defaults.items():
        if field not in st.session_state:
            st.session_state[field] = default

    with st.form("manual_add"):
        manual_title = st.text_input("Problem title", key="manual_title")
        manual_description = st.text_area("Problem description", key="manual_description")
        manual_solution = st.text_area("Verified solution steps", height=220, key="manual_solution")
        manual_keywords = st.text_input("Keywords (optional)", key="manual_keywords")
        manual_notes = st.text_area("Verification notes / evidence", key="manual_notes")
        submitted = st.form_submit_button("Save verified solution")

    if submitted:
        if not all(normalize(x) for x in (manual_title, manual_description, manual_solution)):
            st.error("Title, description, and solution are required.")
        elif len(normalize(manual_notes)) < 10:
            st.error("Add verification notes (at least 10 characters).")
        else:
            try:
                pid, path = add_entry(
                    manual_title,
                    manual_description,
                    manual_solution,
                    "Verified",
                    manual_notes,
                    manual_keywords,
                )
                st.session_state["manual_save_message"] = f"Saved {pid}: {path.name}"
                for field in manual_defaults:
                    st.session_state[field] = ""
                st.rerun()
            except Exception as exc:
                st.error(f"Could not save the solution: {exc}")

    if st.session_state.get("manual_save_message"):
        st.success(st.session_state.pop("manual_save_message"))